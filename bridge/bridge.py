import json
import os
import shutil
import subprocess
import sys
import threading
import time
from datetime import datetime
from http.server import HTTPServer, BaseHTTPRequestHandler
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

PRIMARY_EXCEL_PATH = r"C:\Users\ASUS\Job_Tracker\Job_Applications_Tracker.xlsx"
PENDING_QUEUE_PATH = r"C:\Users\ASUS\Job_Tracker\pending_jobs.json"
BACKUP_EXCEL_DIR = r"E:\JOB APPLICATION"
BACKUP_EXCEL_PATH = r"E:\JOB APPLICATION\Job_Applications_Tracker.xlsx"
PORT = 5252
EXCEL_EXE = r"C:\Program Files\Microsoft Office\Root\Office16\EXCEL.EXE"

TARGET_URL_UPDATES = {
    "APP-001": "https://www.petronas.com/careers",
    "APP-002": "https://maybankjobs.com/#/home",
    "APP-003": "https://www.cimb.com/en/careers/fresh-graduate/the-complete-banker.html",
    "APP-005": "https://www.bnm.gov.my/careers/vacancy",
    "APP-007": "https://maxine.wd3.myworkdayjobs.com/Maxis-Early-Careers"
}

STATUS_COLORS = {
    "Applied": "DCFCE7",        # Soft Green
    "Under Review": "E0F2FE",   # Soft Sky Blue
    "Assessment / Test": "EDE9FE", # Light Purple
    "Interviewing": "FEF3C7",   # Warm Amber
    "Shortlisted": "D1FAE5",    # Soft Teal / Green
    "Offer Received": "86EFAC", # Vibrant Mint
    "Rejected": "FEE2E2",       # Soft Red
    "Target / Backlog": "F8FAFC" # Neutral Light Slate
}

THIN_BORDER = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

def launch_excel():
    # 0. Flush any pending jobs into Excel file first
    flush_pending_to_excel()

    # 1. Terminate any headless zombie Excel processes that block interactive UI
    try:
        subprocess.run(
            ["powershell", "-NoProfile", "-Command", "Get-Process excel -ErrorAction SilentlyContinue | Where-Object { $_.MainWindowTitle -eq '' } | Stop-Process -Force"],
            capture_output=True,
            timeout=3
        )
    except Exception:
        pass

    # 2. Try launching via file association and explorer
    try:
        os.startfile(PRIMARY_EXCEL_PATH)
    except Exception:
        pass

    # 3. Open File Explorer highlighting the file so user has instant physical access
    try:
        subprocess.Popen(["explorer.exe", f"/select,{PRIMARY_EXCEL_PATH}"])
    except Exception:
        pass

    return True, "Excel opening command triggered and file location highlighted in File Explorer!"

def format_and_append_row(ws, job, app_id):
    date_applied = job.get("date_applied") or datetime.now().strftime("%Y-%m-%d")
    company = job.get("company", "").strip() or "N/A"
    job_title = job.get("job_title", "").strip() or "N/A"
    job_type = job.get("job_type", "").strip() or "Graduate / Permanent"
    location = job.get("location", "").strip() or "Malaysia"
    platform = job.get("platform", "").strip() or "Direct"
    status = job.get("status", "").strip() or "Applied"
    url = job.get("url", "").strip()
    email = job.get("email", "").strip()
    salary = job.get("salary", "").strip()
    follow_up = job.get("follow_up", "").strip() or datetime.now().strftime("%Y-%m-%d")
    notes = job.get("notes", "").strip()

    new_row = [
        app_id,
        date_applied,
        company,
        job_title,
        job_type,
        location,
        platform,
        status,
        url,
        email,
        salary,
        follow_up,
        notes
    ]

    ws.append(new_row)
    new_row_idx = ws.max_row
    ws.row_dimensions[new_row_idx].height = 22

    status_bg = STATUS_COLORS.get(status, "FFFFFF")

    for col_idx in range(1, len(new_row) + 1):
        cell = ws.cell(row=new_row_idx, column=col_idx)
        cell.font = Font(name="Calibri", size=9.5)
        cell.border = THIN_BORDER
        
        if col_idx in [1, 2, 5, 8, 12]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx in [11]:
            cell.alignment = Alignment(horizontal="right", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")
            
        if col_idx == 8:
            cell.fill = PatternFill(start_color=status_bg, end_color=status_bg, fill_type="solid")
            cell.font = Font(name="Calibri", size=9.5, bold=True)

    return new_row_idx

def flush_pending_to_excel():
    pending = []
    if os.path.exists(PENDING_QUEUE_PATH):
        try:
            with open(PENDING_QUEUE_PATH, "r", encoding="utf-8") as f:
                pending = json.load(f)
        except Exception:
            pending = []

    if not os.path.exists(PRIMARY_EXCEL_PATH):
        return False

    try:
        wb = openpyxl.load_workbook(PRIMARY_EXCEL_PATH)
        ws = wb.active

        # Update target links if needed
        for r in range(2, ws.max_row + 1):
            aid = str(ws.cell(row=r, column=1).value or "").strip()
            if aid in TARGET_URL_UPDATES:
                ws.cell(row=r, column=9).value = TARGET_URL_UPDATES[aid]

        # Flush pending jobs
        if pending:
            for job in pending:
                app_id = job.get("app_id") or get_next_app_id()
                format_and_append_row(ws, job, app_id)

        wb.save(PRIMARY_EXCEL_PATH)
        wb.close()

        if pending and os.path.exists(PENDING_QUEUE_PATH):
            try:
                os.remove(PENDING_QUEUE_PATH)
            except Exception:
                pass

        if os.path.exists(BACKUP_EXCEL_DIR):
            try:
                shutil.copy2(PRIMARY_EXCEL_PATH, BACKUP_EXCEL_PATH)
            except Exception:
                pass

        return True
    except PermissionError:
        # File is locked by open Excel instance, will flush when closed
        return False
    except Exception as e:
        print(f"Flush notice: {e}")
        return False

def update_job_status(app_id, new_status):
    # Try flushing first in case Excel was just closed
    flush_pending_to_excel()

    # If item is in pending_jobs.json, update it there
    if os.path.exists(PENDING_QUEUE_PATH):
        try:
            with open(PENDING_QUEUE_PATH, "r", encoding="utf-8") as f:
                pending = json.load(f)
            updated_p = False
            for j in pending:
                if str(j.get("app_id", "")).strip() == app_id:
                    j["status"] = new_status
                    updated_p = True
            if updated_p:
                with open(PENDING_QUEUE_PATH, "w", encoding="utf-8") as f:
                    json.dump(pending, f, indent=2)
                return True, f"Updated {app_id} status to '{new_status}' (in live queue)"
        except Exception:
            pass

    # Update in primary Excel file
    if os.path.exists(PRIMARY_EXCEL_PATH):
        try:
            wb = openpyxl.load_workbook(PRIMARY_EXCEL_PATH)
            ws = wb.active
            found = False
            for r in range(2, ws.max_row + 1):
                if str(ws.cell(row=r, column=1).value or "").strip() == app_id:
                    status_cell = ws.cell(row=r, column=8)
                    status_cell.value = new_status
                    status_bg = STATUS_COLORS.get(new_status, "FFFFFF")
                    status_cell.fill = PatternFill(start_color=status_bg, end_color=status_bg, fill_type="solid")
                    status_cell.font = Font(name="Calibri", size=9.5, bold=True)
                    found = True
                    break
            if found:
                wb.save(PRIMARY_EXCEL_PATH)
                wb.close()
                if os.path.exists(BACKUP_EXCEL_DIR):
                    try:
                        shutil.copy2(PRIMARY_EXCEL_PATH, BACKUP_EXCEL_PATH)
                    except Exception:
                        pass
                return True, f"Updated {app_id} status to '{new_status}' in Excel!"
            else:
                wb.close()
                return False, f"Application {app_id} not found in Excel"
        except PermissionError:
            return False, "Excel is currently open on your desktop. Please close Excel to update from browser, or edit directly in Excel!"
        except Exception as e:
            return False, f"Error updating status: {e}"

    return False, "Tracker Excel file not found"

def get_next_app_id():
    max_id_num = 0
    if os.path.exists(PRIMARY_EXCEL_PATH):
        try:
            wb = openpyxl.load_workbook(PRIMARY_EXCEL_PATH, data_only=True)
            ws = wb.active
            for row in range(2, ws.max_row + 1):
                cell_val = str(ws.cell(row=row, column=1).value or "")
                if cell_val.startswith("APP-"):
                    try:
                        num = int(cell_val.replace("APP-", ""))
                        if num > max_id_num:
                            max_id_num = num
                    except ValueError:
                        pass
            wb.close()
        except Exception:
            pass

    if os.path.exists(PENDING_QUEUE_PATH):
        try:
            with open(PENDING_QUEUE_PATH, "r", encoding="utf-8") as f:
                pending = json.load(f)
            for j in pending:
                aid = str(j.get("app_id", ""))
                if aid.startswith("APP-"):
                    try:
                        num = int(aid.replace("APP-", ""))
                        if num > max_id_num:
                            max_id_num = num
                    except ValueError:
                        pass
        except Exception:
            pass

    return f"APP-{max_id_num + 1:03d}"

def append_job_to_excel(job):
    # Try flushing first in case Excel closed
    flush_pending_to_excel()

    next_id = get_next_app_id()
    job["app_id"] = next_id

    try:
        wb = openpyxl.load_workbook(PRIMARY_EXCEL_PATH)
        ws = wb.active
        new_row_idx = format_and_append_row(ws, job, next_id)
        wb.save(PRIMARY_EXCEL_PATH)
        wb.close()

        if os.path.exists(BACKUP_EXCEL_DIR):
            try:
                shutil.copy2(PRIMARY_EXCEL_PATH, BACKUP_EXCEL_PATH)
            except Exception:
                pass

        return next_id, new_row_idx - 1, "saved"
    except PermissionError:
        # Excel is currently open on desktop! Queue it safely!
        pending = []
        if os.path.exists(PENDING_QUEUE_PATH):
            try:
                with open(PENDING_QUEUE_PATH, "r", encoding="utf-8") as f:
                    pending = json.load(f)
            except Exception:
                pending = []
        pending.append(job)
        with open(PENDING_QUEUE_PATH, "w", encoding="utf-8") as f:
            json.dump(pending, f, indent=2)

        return next_id, len(pending), "queued"

def get_tracker_html():
    # Attempt a flush first
    flush_pending_to_excel()

    rows = []
    if os.path.exists(PRIMARY_EXCEL_PATH):
        try:
            wb = openpyxl.load_workbook(PRIMARY_EXCEL_PATH, data_only=True)
            ws = wb.active
            for r in range(1, ws.max_row + 1):
                row_vals = [str(ws.cell(row=r, column=c).value or '') for c in range(1, ws.max_column + 1)]
                rows.append(row_vals)
            wb.close()
        except Exception as e:
            rows = [[f"Error loading Excel: {e}"]]

    # Also check if any pending jobs in queue
    pending_jobs = []
    if os.path.exists(PENDING_QUEUE_PATH):
        try:
            with open(PENDING_QUEUE_PATH, "r", encoding="utf-8") as f:
                pending_jobs = json.load(f)
        except Exception:
            pending_jobs = []

    html = """<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <title>Job Applications Tracker | Irfan Firdaus</title>
  <link rel="preconnect" href="https://fonts.googleapis.com">
  <link href="https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700;800&family=JetBrains+Mono:wght@400;500;600&display=swap" rel="stylesheet">
  <style>
    * { box-sizing: border-box; margin: 0; padding: 0; font-family: 'Plus Jakarta Sans', sans-serif; }
    body { background: #0B1120; color: #F1F5F9; padding: 24px; font-size: 13px; }
    .header { display: flex; justify-content: space-between; align-items: center; margin-bottom: 20px; border-bottom: 1px solid #1E293B; padding-bottom: 16px; }
    h1 { font-size: 20px; font-weight: 800; color: #FFFFFF; display: flex; align-items: center; gap: 10px; }
    .badge { background: #1E3A8A; color: #93C5FD; font-size: 11px; padding: 3px 8px; border-radius: 12px; font-weight: 600; }
    .btn { background: #10B981; color: white; padding: 8px 14px; border-radius: 8px; text-decoration: none; font-weight: 600; font-size: 12px; display: inline-flex; align-items: center; gap: 6px; cursor: pointer; border: none; }
    .btn:hover { background: #059669; }
    .btn-secondary { background: #1E293B; color: #CBD5E1; }
    .btn-secondary:hover { background: #334155; }
    .card { background: #0F172A; border: 1px solid #1E293B; border-radius: 12px; overflow: hidden; box-shadow: 0 4px 6px -1px rgba(0,0,0,0.3); }
    table { width: 100%; border-collapse: collapse; text-align: left; }
    th { background: #0F2C59; color: #FFFFFF; font-weight: 700; padding: 10px 14px; text-transform: uppercase; font-size: 10.5px; letter-spacing: 0.5px; border-bottom: 1px solid #1E293B; }
    td { padding: 9px 14px; border-bottom: 1px solid #1E293B; font-size: 12px; color: #CBD5E1; }
    tr:hover td { background: #1E293B80; }
    .mono { font-family: 'JetBrains Mono', monospace; font-size: 11.5px; }

    /* Base Status Dropdown */
    .status-select { 
      border-radius: 8px; 
      padding: 5px 10px; 
      font-size: 11.5px; 
      font-weight: 700; 
      cursor: pointer; 
      outline: none; 
      transition: all 0.2s ease;
      display: inline-block;
      box-shadow: 0 1px 3px rgba(0,0,0,0.3);
    }
    .status-select option { 
      background: #0F172A; 
      color: #F1F5F9; 
      font-weight: 600; 
      padding: 6px; 
    }

    /* Status Colors */
    /* 1. Applied - Fresh Emerald Green */
    .status-style-Applied { 
      background: rgba(16, 185, 129, 0.2) !important; 
      color: #34D399 !important; 
      border: 1.5px solid #10B981 !important; 
    }

    /* 2. Rejected - Clear Red */
    .status-style-Rejected { 
      background: rgba(239, 68, 68, 0.22) !important; 
      color: #F87171 !important; 
      border: 1.5px solid #EF4444 !important; 
    }

    /* 3. Assessment / Test - Distinctive Purple */
    .status-style-Assessment---Test, .status-style-Assessment-Test { 
      background: rgba(168, 85, 247, 0.22) !important; 
      color: #C084FC !important; 
      border: 1.5px solid #A855F7 !important; 
    }

    /* 4. Under Review - Sky Blue */
    .status-style-Under-Review { 
      background: rgba(14, 165, 233, 0.2) !important; 
      color: #38BDF8 !important; 
      border: 1.5px solid #0EA5E9 !important; 
    }

    /* 5. Shortlisted - Bright Teal */
    .status-style-Shortlisted { 
      background: rgba(20, 184, 166, 0.22) !important; 
      color: #2DD4BF !important; 
      border: 1.5px solid #14B8A6 !important; 
    }

    /* 6. Interviewing - Warm Amber */
    .status-style-Interviewing { 
      background: rgba(245, 158, 11, 0.22) !important; 
      color: #FBBF24 !important; 
      border: 1.5px solid #F59E0B !important; 
    }

    /* 7. Offer Received - Mint Green */
    .status-style-Offer-Received { 
      background: rgba(34, 197, 94, 0.26) !important; 
      color: #86EFAC !important; 
      border: 1.5px solid #22C55E !important; 
    }

    /* 8. Target / Backlog & Neutral - Clean Dark Slate (No Distraction) */
    .status-style-Target---Backlog, .status-style-default { 
      background: #1E293B !important; 
      color: #94A3B8 !important; 
      border: 1px solid #334155 !important; 
    }
  </style>
</head>
<body>
  <div class="header">
    <div>
      <h1><span>📊 Job Applications Tracker</span> <span class="badge">Live View</span></h1>
      <p style="color: #64748B; font-size: 12px; margin-top: 4px;">Muhammad Irfan Firdaus • C:\\Users\\ASUS\\Job_Tracker\\Job_Applications_Tracker.xlsx</p>
    </div>
    <div style="display: flex; gap: 8px; align-items: center; flex-wrap: wrap;">
      <a href="/download-excel" class="btn" style="background: #2563EB; text-decoration: none;">📥 Download / Open .xlsx</a>
      <button onclick="launchExcel(this)" class="btn" style="cursor: pointer; border: none;">📊 Open in Microsoft Excel</button>
      <button onclick="location.reload()" class="btn btn-secondary">🔄 Refresh</button>
    </div>
  </div>

  <div id="launchToast" style="display: none; padding: 10px 14px; background: #064E3B; border: 1px solid #059669; color: #A7F3D0; border-radius: 8px; margin-bottom: 16px; font-weight: 600; font-size: 12px;"></div>
"""
    if pending_jobs:
        html += f"""  <div style="background: rgba(245, 158, 11, 0.15); border: 1px solid #F59E0B; padding: 10px 14px; border-radius: 8px; margin-bottom: 16px; color: #FDE68A; font-size: 12px;">
    ⚠️ <b>{len(pending_jobs)} new application(s) queued live!</b> Because Excel is open on your desktop, they will automatically write to your .xlsx file the moment Excel is closed.
  </div>"""

    html += """  <div class="card">
    <table>
      <thead>
        <tr>"""
    if rows:
        headers = rows[0]
        for h in headers[:9]:
            html += f"<th>{h}</th>"
        html += "<th>Salary</th><th>Notes</th></tr></thead><tbody>"

        STATUS_OPTIONS = [
            "Applied",
            "Under Review",
            "Assessment / Test",
            "Interviewing",
            "Shortlisted",
            "Offer Received",
            "Target / Backlog",
            "Rejected"
        ]

        def get_status_class(status):
            clean = (status or "").replace(" / ", "---").replace(" ", "-")
            return f"status-style-{clean}"

        for row in rows[1:]:
            html += "<tr>"
            # App ID
            aid = row[0] if len(row) > 0 else ''
            html += f"<td class='mono' style='color: #60A5FA; font-weight: bold;'>{aid}</td>"
            # Date
            html += f"<td class='mono' style='color: #94A3B8;'>{row[1] if len(row) > 1 else ''}</td>"
            # Company
            html += f"<td style='color: #FFFFFF; font-weight: 600;'>{row[2] if len(row) > 2 else ''}</td>"
            # Job Title
            html += f"<td style='color: #E2E8F0;'>{row[3] if len(row) > 3 else ''}</td>"
            # Job Type
            html += f"<td style='color: #94A3B8;'>{row[4] if len(row) > 4 else ''}</td>"
            # Location
            html += f"<td style='color: #94A3B8;'>{row[5] if len(row) > 5 else ''}</td>"
            # Platform
            html += f"<td style='color: #CBD5E1;'>{row[6] if len(row) > 6 else ''}</td>"
            # Status
            status_text = row[7] if len(row) > 7 else 'Applied'
            s_class = get_status_class(status_text)
            opts = "".join([f"<option value='{s}' {'selected' if s == status_text else ''}>{s}</option>" for s in STATUS_OPTIONS])
            html += f"<td><select class='status-select {s_class}' onchange=\"updateStatus(this, '{aid}')\">{opts}</select></td>"
            # URL
            url_text = row[8] if len(row) > 8 else ''
            if url_text.startswith("http"):
                html += f"<td><a href='{url_text}' target='_blank' style='color: #38BDF8;'>Link</a></td>"
            else:
                html += f"<td style='color: #64748B;'>-</td>"
            # Salary
            sal = row[10] if len(row) > 10 else ''
            html += f"<td class='mono' style='color: #34D399; font-weight: 600;'>{sal}</td>"
            # Notes
            notes = row[12] if len(row) > 12 else ''
            html += f"<td style='color: #94A3B8; max-width: 320px; font-size: 11px;'>{notes}</td>"
            html += "</tr>"

        # Also render pending queued jobs in live web view
        for pj in pending_jobs:
            status_text = pj.get("status", "Applied")
            s_class = get_status_class(status_text)
            aid = pj.get('app_id', 'NEW')
            opts = "".join([f"<option value='{s}' {'selected' if s == status_text else ''}>{s}</option>" for s in STATUS_OPTIONS])
            url_text = pj.get("url", "")
            link_html = f"<a href='{url_text}' target='_blank' style='color: #38BDF8;'>Link</a>" if url_text.startswith("http") else "<td style='color: #64748B;'>-</td>"
            html += f"""<tr style="background: rgba(16, 185, 129, 0.08);">
              <td class='mono' style='color: #34D399; font-weight: bold;'>{aid} <span style="font-size: 9px; background: #059669; color: white; padding: 1px 4px; border-radius: 4px;">LIVE</span></td>
              <td class='mono' style='color: #94A3B8;'>{pj.get('date_applied', '')}</td>
              <td style='color: #FFFFFF; font-weight: 600;'>{pj.get('company', '')}</td>
              <td style='color: #E2E8F0;'>{pj.get('job_title', '')}</td>
              <td style='color: #94A3B8;'>{pj.get('job_type', '')}</td>
              <td style='color: #94A3B8;'>{pj.get('location', '')}</td>
              <td style='color: #CBD5E1;'>{pj.get('platform', '')}</td>
              <td><select class='status-select {s_class}' onchange=\"updateStatus(this, '{aid}')\">{opts}</select></td>
              <td>{link_html}</td>
              <td class='mono' style='color: #34D399; font-weight: 600;'>{pj.get('salary', '')}</td>
              <td style='color: #94A3B8; max-width: 320px; font-size: 11px;'>{pj.get('notes', '')}</td>
            </tr>"""

        html += """</tbody></table></div>
  <script>
    function updateSelectColor(selectEl) {
      const val = selectEl.value;
      const clean = (val || "").replace(/ \/ /g, '---').replace(/ /g, '-');
      const classes = Array.from(selectEl.classList).filter(c => !c.startsWith('status-style-'));
      selectEl.className = classes.join(' ');
      selectEl.classList.add('status-style-' + clean);
    }

    async function launchExcel(btn) {
      const orig = btn.innerHTML;
      btn.innerHTML = "⏳ Opening...";
      btn.disabled = true;
      const toast = document.getElementById("launchToast");
      try {
        const res = await fetch("/open-excel");
        const data = await res.json();
        
        // Also trigger browser download/open directly in user's active window
        const iframe = document.createElement("iframe");
        iframe.style.display = "none";
        iframe.src = "/download-excel";
        document.body.appendChild(iframe);
        setTimeout(() => iframe.remove(), 10000);

        toast.style.display = "block";
        toast.style.background = "#064E3B";
        toast.style.color = "#A7F3D0";
        toast.innerHTML = "✅ Excel opening command triggered! A copy is also downloading in Opera. Check your taskbar / downloads, or use the Desktop shortcut: <code>Job Applications Tracker.lnk</code>";
        setTimeout(() => { toast.style.display = "none"; }, 10000);
      } catch (err) {
        toast.style.display = "block";
        toast.style.background = "#7F1D1D";
        toast.style.color = "#FECACA";
        toast.innerText = "Error: " + err;
      } finally {
        btn.innerHTML = orig;
        btn.disabled = false;
      }
    }

    async function updateStatus(selectEl, appId) {
      const newStatus = selectEl.value;
      updateSelectColor(selectEl);
      const toast = document.getElementById("launchToast");
      try {
        const res = await fetch("/update-status", {
          method: "POST",
          headers: { "Content-Type": "application/json" },
          body: JSON.stringify({ app_id: appId, status: newStatus })
        });
        const data = await res.json();
        toast.style.display = "block";
        if (data.success) {
          toast.style.background = "#064E3B";
          toast.style.color = "#A7F3D0";
          toast.innerHTML = "✅ " + data.message;
        } else {
          toast.style.background = "#7F1D1D";
          toast.style.color = "#FECACA";
          toast.innerHTML = "⚠️ " + data.message;
        }
        setTimeout(() => { toast.style.display = "none"; }, 5000);
      } catch (err) {
        toast.style.display = "block";
        toast.style.background = "#7F1D1D";
        toast.style.color = "#FECACA";
        toast.innerText = "Error updating status: " + err;
      }
    }
  </script>
</body></html>"""
    return html

class TrackerHandler(BaseHTTPRequestHandler):
    def _send_cors_headers(self):
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")

    def do_OPTIONS(self):
        self.send_response(200)
        self._send_cors_headers()
        self.end_headers()

    def do_GET(self):
        if self.path == "/status" or self.path == "/":
            total_records = 0
            if os.path.exists(PRIMARY_EXCEL_PATH):
                try:
                    wb = openpyxl.load_workbook(PRIMARY_EXCEL_PATH, read_only=True)
                    ws = wb.active
                    total_records = max(0, ws.max_row - 1)
                    wb.close()
                except Exception:
                    pass

            pending_count = 0
            if os.path.exists(PENDING_QUEUE_PATH):
                try:
                    with open(PENDING_QUEUE_PATH, "r", encoding="utf-8") as f:
                        pending_count = len(json.load(f))
                except Exception:
                    pass

            external_connected = os.path.exists(BACKUP_EXCEL_DIR)

            response_data = {
                "status": "connected",
                "excel_path": PRIMARY_EXCEL_PATH,
                "external_backup_available": external_connected,
                "total_records": total_records + pending_count,
                "pending_queue": pending_count,
                "message": "Job Tracker Sync Bridge is Active"
            }
            self.send_response(200)
            self._send_cors_headers()
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(json.dumps(response_data).encode("utf-8"))

        elif self.path == "/view":
            html_content = get_tracker_html()
            self.send_response(200)
            self._send_cors_headers()
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.end_headers()
            self.wfile.write(html_content.encode("utf-8"))

        elif self.path == "/download-excel":
            flush_pending_to_excel()
            if os.path.exists(PRIMARY_EXCEL_PATH):
                try:
                    with open(PRIMARY_EXCEL_PATH, "rb") as f:
                        file_bytes = f.read()
                    self.send_response(200)
                    self._send_cors_headers()
                    self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                    self.send_header("Content-Disposition", 'attachment; filename="Job_Applications_Tracker.xlsx"')
                    self.send_header("Content-Length", str(len(file_bytes)))
                    self.end_headers()
                    self.wfile.write(file_bytes)
                except Exception as e:
                    self.send_response(500)
                    self._send_cors_headers()
                    self.end_headers()
                    self.wfile.write(f"Error streaming file: {e}".encode("utf-8"))
            else:
                self.send_response(404)
                self.end_headers()

        elif self.path == "/open-excel":
            success, msg = launch_excel()
            self.send_response(200 if success else 500)
            self._send_cors_headers()
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(json.dumps({"success": success, "message": msg}).encode("utf-8"))

        else:
            self.send_response(404)
            self._send_cors_headers()
            self.end_headers()

    def do_POST(self):
        if self.path == "/add-job":
            content_length = int(self.headers.get("Content-Length", 0))
            post_data = self.rfile.read(content_length)
            try:
                job_data = json.loads(post_data.decode("utf-8"))
                app_id, total, status_type = append_job_to_excel(job_data)
                
                msg = f"Successfully added {app_id} to Excel!" if status_type == "saved" else f"Added {app_id} to live tracker! (Will sync to Excel file as soon as Excel closes)"

                response_data = {
                    "success": True,
                    "app_id": app_id,
                    "total_records": total,
                    "status_type": status_type,
                    "message": msg
                }
                self.send_response(200)
                self._send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps(response_data).encode("utf-8"))
            except Exception as e:
                self.send_response(500)
                self._send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps({"success": False, "error": str(e)}).encode("utf-8"))

        elif self.path == "/update-status":
            content_length = int(self.headers.get("Content-Length", 0))
            post_data = self.rfile.read(content_length)
            try:
                payload = json.loads(post_data.decode("utf-8"))
                app_id = payload.get("app_id", "").strip()
                status = payload.get("status", "").strip()
                success, msg = update_job_status(app_id, status)
                self.send_response(200 if success else 400)
                self._send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps({"success": success, "message": msg}).encode("utf-8"))
            except Exception as e:
                self.send_response(500)
                self._send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps({"success": False, "message": str(e)}).encode("utf-8"))

        else:
            self.send_response(404)
            self._send_cors_headers()
            self.end_headers()

    def log_message(self, format, *args):
        return

def background_worker():
    while True:
        time.sleep(3)
        flush_pending_to_excel()

def run_server():
    server_address = ("127.0.0.1", PORT)
    httpd = HTTPServer(server_address, TrackerHandler)
    print(f"=====================================================")
    print(f"   JOB TRACKER EXCEL SYNC BRIDGE (C: DRIVE MASTER)")
    print(f"   Primary Excel: {PRIMARY_EXCEL_PATH}")
    print(f"   Web Dashboard: http://127.0.0.1:{PORT}/view")
    print(f"=====================================================")
    
    # Start auto-flush worker in background
    t = threading.Thread(target=background_worker, daemon=True)
    t.start()

    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        httpd.server_close()

if __name__ == "__main__":
    run_server()
